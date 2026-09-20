import csv
import io
import json
from pathlib import Path


class UnsupportedFileTypeError(Exception):
    pass


class DocumentParseError(Exception):
    pass


SUPPORTED_EXTENSIONS = {
    "txt", "md", "pdf", "docx", "csv",
    "xlsx", "pptx", "json", "html", "htm", "xml"
}


def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = Path(filename or "").suffix.lower().lstrip(".")

    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '.{ext}'. Supported types: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    try:
        if ext == "txt":
            return _extract_txt(file_bytes)

        if ext == "md":
            return _extract_txt(file_bytes)

        if ext == "pdf":
            return _extract_pdf(file_bytes)

        if ext == "docx":
            return _extract_docx(file_bytes)

        if ext == "csv":
            return _extract_csv(file_bytes)

        if ext == "xlsx":
            return _extract_xlsx(file_bytes)

        if ext == "pptx":
            return _extract_pptx(file_bytes)

        if ext == "json":
            return _extract_json(file_bytes)

        if ext in ("html", "htm"):
            return _extract_html(file_bytes)

        if ext == "xml":
            return _extract_txt(file_bytes)

    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(
            f"Could not read {ext.upper()} file: {e}"
        )

    raise UnsupportedFileTypeError(f"Unsupported file type: .{ext}")


def _extract_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace").strip()


def _extract_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(file_bytes))

        text_parts = [
            page.extract_text() or ""
            for page in reader.pages
        ]

        text = "\n".join(text_parts).strip()

        if not text:
            raise DocumentParseError(
                "PDF contains no extractable text. "
                "Scanned/image-only PDFs require OCR."
            )

        return text

    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"Could not read PDF: {e}")


def _extract_docx(file_bytes: bytes) -> str:
    try:
        import zipfile
        import xml.etree.ElementTree as ET
        from docx import Document

        buffer = io.BytesIO(file_bytes)

        # A real DOCX is an OOXML ZIP package. Renamed .doc/.pdf files
        # must still be rejected rather than producing garbage text.
        if not zipfile.is_zipfile(buffer):
            raise DocumentParseError(
                "The uploaded file is not a valid DOCX file. "
                "Please upload a real .docx Word document."
            )

        buffer.seek(0)
        doc = Document(buffer)

        parts = []

        for paragraph in doc.paragraphs:
            value = paragraph.text.strip()
            if value:
                parts.append(value)

        for table in doc.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text and cell.text.strip()
                ]
                if values:
                    parts.append(" | ".join(values))

        text = "\n".join(parts).strip()

        if text:
            return text

        # python-docx does not expose every OOXML text container (for
        # example some text boxes/shapes, headers/footers and notes).
        # Inspect Word XML text nodes before declaring the file empty.
        xml_parts = []

        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            candidate_names = [
                name
                for name in archive.namelist()
                if name.startswith("word/")
                and name.endswith(".xml")
                and (
                    name == "word/document.xml"
                    or name.startswith("word/header")
                    or name.startswith("word/footer")
                    or name in {
                        "word/footnotes.xml",
                        "word/endnotes.xml",
                        "word/comments.xml",
                    }
                )
            ]

            for name in candidate_names:
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError:
                    continue

                values = []
                for node in root.iter():
                    if node.tag.endswith("}t") and node.text:
                        value = node.text.strip()
                        if value:
                            values.append(value)

                if values:
                    xml_parts.append(" ".join(values))

        text = "\n".join(xml_parts).strip()

        if not text:
            raise DocumentParseError(
                "DOCX contains no extractable text. "
                "If the document contains only images/scans, OCR is required."
            )

        return text

    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"Could not read DOCX: {e}")


def _extract_csv(file_bytes: bytes) -> str:
    try:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)

        if not rows:
            return ""

        header = rows[0]
        lines = []

        for row in rows[1:]:
            pairs = [
                f"{h}: {v}"
                for h, v in zip(header, row)
            ]
            lines.append(", ".join(pairs))

        return "\n".join(lines).strip()

    except Exception as e:
        raise DocumentParseError(f"Could not read CSV: {e}")


def _extract_xlsx(file_bytes: bytes) -> str:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True
        )

        lines = []

        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value) if value is not None else ""
                    for value in row
                ]

                if any(values):
                    lines.append(" | ".join(values))

        return "\n".join(lines).strip()

    except Exception as e:
        raise DocumentParseError(f"Could not read XLSX: {e}")


def _extract_pptx(file_bytes: bytes) -> str:
    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(file_bytes))

        slides = []

        for index, slide in enumerate(
            presentation.slides,
            start=1
        ):
            slides.append(f"Slide {index}")

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slides.append(shape.text)

        return "\n".join(slides).strip()

    except Exception as e:
        raise DocumentParseError(f"Could not read PPTX: {e}")


def _extract_json(file_bytes: bytes) -> str:
    try:
        data = json.loads(
            file_bytes.decode("utf-8-sig")
        )

        return json.dumps(
            data,
            indent=2,
            ensure_ascii=False
        )

    except Exception as e:
        raise DocumentParseError(f"Could not read JSON: {e}")


def _extract_html(file_bytes: bytes) -> str:
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(
            file_bytes.decode(
                "utf-8",
                errors="replace"
            ),
            "html.parser"
        )

        return soup.get_text(
            separator="\n",
            strip=True
        )

    except Exception as e:
        raise DocumentParseError(f"Could not read HTML: {e}")